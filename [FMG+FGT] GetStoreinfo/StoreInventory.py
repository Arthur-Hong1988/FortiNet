# Edit by Liam
# This script is to get all devices that each store has.

from FMGRegisteredDevices import Getdeviceinfo
from pprint import pprint
from netmiko import ConnectHandler
from getpass import getpass
from netmiko.exceptions import (NetmikoTimeoutException,NetmikoAuthenticationException)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path


# When you enter the storecode, you will get the FGT's IP address
# This address is used for Netmiko
def get_device_ip(storecode):
    alldevices = Getdeviceinfo()
    for device in alldevices:
        if device["hostname"] == storecode:
            return device["ip"],device["adom"],device["hostname"],device["sn"],device["platform"]

    print("Please enter correct storecode")
    return None

# Use netmiko to connect to FortiGate and execute some command to get information
def connect_fgt(ip,username,password):
    device = {
        "device_type": "fortinet",
        "host": ip,
        "username": username,
        "password": password,
        "port": 22,
    }

    try:
        with ConnectHandler(**device) as connection:
            output_sw = connection.send_command(
                "execute switch-controller get-conn-status",
                read_timeout=30
            )
            output_ap = connection.send_command(
                "show wireless-controller wtp",
                read_timeout=30
            )
            return output_sw, output_ap

    except NetmikoAuthenticationException:
        print(f"{ip}: Authentication failed")

    except NetmikoTimeoutException:
        print(f"{ip}: Connection timeout")

    except Exception as e:
        print(f"{ip}: Error: {e}")

    return None

def get_model_from_sn(serial):
    MODEL_MAP = {
        "FP231K": "FortiAP-231K",
        "PU431F": "FortiAP-U431F",
        "FP431F": "FortiAP-431F",
        "FP321E": "FortiAP-321E",
        "FP431G": "FortiAP-431G",
        "S108EF": "FortiSwitch-108E-FPOE",
        "S108FF": "FortiSwitch-108F-FPOE",
        "S124GF": "FortiSwitch-124G-FPOE",
        "S124FF": "FortiSwitch-124F-FPOE",
        "S124EF": "FortiSwitch-124E-FPOE",
        "S148FF": "FortiSwitch-148F-FPOE",

    }
    for prefix, model in MODEL_MAP.items():
        if serial.startswith(prefix):
            return model
    return "Unknown"

# This function is for getting infomation of FortiSwitch
def get_sw(output):
    swlist = []
    if not output:
        return swlist

    for line in output.splitlines():
        if "Authorized/Up" in line:
            parts = line.split()
            switch_name = parts[0]
            serial = parts[-1]

            swlist.append({
                "hostname": switch_name,
                "sn": serial,
                "platform": get_model_from_sn(serial)
            })
    return swlist


def get_ap(output):
    aplist = []
    serial = None
    if not output:
        return aplist

    for line in output.splitlines():
        line = line.strip()
        if line.startswith("edit "):
            serial = line.split('"')[1]
        elif line.startswith("set name ") and serial:       # ensure already match the serial
            name = line.split('"')[1]
            aplist.append({
                "hostname": name,
                "sn": serial,
                "platform": get_model_from_sn(serial)
            })
            serial = None
    return aplist


def write_storeinfo_to_excel(storeinfo):
    # Create a new Excel workbook in memory.
    workbook = Workbook()
    # Get the active worksheet from the workbook.
    worksheet = workbook.active
    worksheet.title = "Device Information"

    # Write the storecode in the first row
    worksheet.append([
        "Storecode",
        storeinfo["storecode"],
        "",
        ""
    ])

    # Write the column headers in the second row
    headers = [
        "Device Type",
        "Hostname",
        "Platform",
        "SN"
    ]
    worksheet.append(headers)

    # Write FortiGate information
    fgt = storeinfo["FGT"]

    worksheet.append([
        "FGT",
        fgt["hostname"],
        fgt["platform"],
        fgt["sn"]
    ])

    # Write FortiSwitch information
    for switch in storeinfo["SW"]:
        worksheet.append([
            "SW",
            switch["hostname"],
            switch["platform"],
            switch["sn"]
        ])

    # Write FortiAP information
    for ap in storeinfo["AP"]:
        worksheet.append([
            "AP",
            ap["hostname"],
            ap["platform"],
            ap["sn"]
        ])

    # Format the storecode row
    worksheet["A1"].font = Font(
        color="FFFFFF",
        bold=True
    )

    worksheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    worksheet["B1"].font = Font(bold=True)

    # Format the header row
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    for cell in worksheet[2]:
        cell.font = Font(
            color="FFFFFF",
            bold=True
        )
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Set column widths
    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 40
    worksheet.column_dimensions["C"].width = 28
    worksheet.column_dimensions["D"].width = 25

    # Freeze the header rows
    worksheet.freeze_panes = "A3"

    # Generate the filename
    filename = (
        f"{storeinfo['storecode']}_Device_Information.xlsx"
    )
    output_path = Path(filename).resolve()
    workbook.save(output_path)
    return output_path


def main():
    storecode = input("Enter storecode: ")
    username = input("Enter your username: ")
    password = getpass("Enter your password: ")
    deviceinfo = get_device_ip(storecode)
    if deviceinfo is None:
        return
    IP,ADOM,hostname,sn,platform = deviceinfo
    result = connect_fgt(IP, username, password)

    if result is None:
        print(f"Failed to connect to FGT: {storecode}")

    else:
        output_sw, output_ap = result
        swlist = get_sw(output_sw)
        aplist = get_ap(output_ap)

        # Combine all store information
        storeinfo = {
            "storecode": storecode,
            "adom": ADOM,
            "FGT": {
                "hostname": hostname,
                "platform": platform,
                "sn": sn,
                "ip": IP
            },
            # 🟢 SW 不存在时，这里就是 []
            "SW": swlist,
            # 🟢 AP 不存在时，这里就是 []
            "AP": aplist
        }
        #pprint(storeinfo)

        output_path = write_storeinfo_to_excel(storeinfo)
        print(f"\nExcel file saved to: {output_path}")


if __name__ == "__main__":
    main()


